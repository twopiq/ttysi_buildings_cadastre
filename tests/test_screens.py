"""Yangi dizayndagi ekranlar: bosh sahifa, muharrir paneli, import va eksport."""

import os

import pytest

from tests.synthetic import ROOMS, make_plan
from ttysi_cadastre.model import Building, Equipment, Floor, Room


@pytest.fixture(scope="module")
def qapp():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _building():
    b = Building(name="Sinov binosi")
    for n in (1, 2):
        f = Floor(name=f"{n}-qavat", width=1000, height=700, scale_m_per_px=0.02)
        f.rooms = [Room(points=[(100, 100), (500, 100), (500, 400), (100, 400)], name=f"{n}01")]
        f.equipment = [Equipment(code="K", x=200, y=200), Equipment(code="T", x=300, y=200, status="planned")]
        b.floors.append(f)
    b.floors[1].equipment.append(Equipment(code="?", x=400, y=300))
    return b


def _window(b=None):
    from ttysi_cadastre.ui.main_window import MainWindow
    w = MainWindow()
    if b is not None:
        w.building = b
        w.refresh_all()
    return w


def _close(w):
    w.dirty = False
    w.close()


def test_home_page_recent_projects(qapp, tmp_path):
    from ttysi_cadastre.ui.home import ProjectCard, forget, recent, remember
    w = _window()
    assert w.pages.currentWidget() is w.home                 # dastur bosh sahifadan boshlanadi
    w.settings.remove("recent/projects")
    b = _building()
    path = str(tmp_path / "bino.ttysi")
    b.save(path)
    w.open_project(path)
    assert w.pages.currentWidget() is w.editor
    items = recent(w.settings)
    assert items[0]["path"] == os.path.abspath(path)
    assert items[0]["floors"] == 2 and items[0]["ready"] == 1 and items[0]["unknown"] == 1
    assert items[0]["counts"] == {"K": 2} and items[0]["planned"] == 2
    remember(w.settings, b, path)                            # takrorlanmaydi
    assert len(recent(w.settings)) == 1
    w.show_home()
    cards = [c for c in w.home.cards if isinstance(c, ProjectCard)]
    assert len(cards) == 1 and cards[0].isEnabled()
    forget(w.settings, path)
    assert recent(w.settings) == []
    _close(w)


def test_editor_tabs_steps_and_tiles(qapp):
    b = _building()
    w = _window(b)
    assert len(w.floor_tabs.buttons) == 2
    assert w.floor_tabs.buttons[0].isChecked()
    assert w.floor_tabs.buttons[1].findChild(type(w.project_name), "badge").text() == "1"
    w.floor_tabs.selected.emit(1)
    assert w.floor is b.floors[1]
    states = [c.property("state") for c in w.steps.chips]
    assert states == ["done", "active", "warn", "todo"]
    assert w.tiles["K"].value.text() == "1"
    assert w.tiles["planned"].value.text() == "+1"
    assert not w.tiles["?"].isHidden()
    assert "tasdiqlash" in w.review_link.text()
    # xona tanlansa - o'ng panelda xona va undagi jihozlar
    room = w.scene.room_items()[0]
    room.setSelected(True)
    assert w.props.currentIndex() == 1
    assert w.props_title.text() == "TANLANGAN XONA"
    texts = [w.room_eq.itemAt(i).widget().text() for i in range(w.room_eq.count()) if w.room_eq.itemAt(i).widget()]
    assert "K. × 1" in texts and "T × 1 kerak" in texts
    _close(w)


def test_export_dialog_and_run(qapp, tmp_path):
    from ttysi_cadastre.ui.export_dialog import ExportDialog
    b = _building()
    w = _window(b)
    d = ExportDialog(w, "svg")
    assert d.kind() == "svg" and d.raster_row.isHidden()
    d.r_all.setChecked(True)
    assert d.floors() == b.floors
    assert not d.warn.isHidden()                             # 2-qavatda "?" bor
    d.c_planned.setChecked(False)
    assert d.options().planned is False
    d.cards["excel"].setChecked(True)
    assert d.scope_box.isHidden() and not d.c_names.isEnabled()
    d.cards["dxf"].setChecked(True)
    assert d.c_planned.isEnabled() and not d.c_names.isEnabled()
    d.reject()

    paths = [str(tmp_path / "1.svg"), str(tmp_path / "2.svg")]
    assert w.run_export("svg", b.floors, d.options(), paths=paths) == paths
    assert all(os.path.getsize(p) > 500 for p in paths)
    pdf = str(tmp_path / "hammasi.pdf")
    w.run_export("pdf", b.floors, d.options(), paths=[pdf])
    assert os.path.getsize(pdf) > 500
    dxf = str(tmp_path / "e.dxf")
    w.run_export("dxf", [b.floors[0]], d.options(), paths=[dxf])
    import ezdxf
    layers = {e.dxf.layer for e in ezdxf.readfile(dxf).modelspace()}
    assert "AT_T_KERAK" not in layers                         # rejadagi jihozlar chiqarilmadi
    _close(w)


def test_export_progress_status(qapp, tmp_path):
    """Eksport holati: tugadi / qisman / xato, qayta urinish."""
    from ttysi_cadastre.render import RenderOptions
    b = _building()
    w = _window(b)
    ok = [str(tmp_path / "a.svg"), str(tmp_path / "b.svg")]
    assert w.run_export("svg", b.floors, RenderOptions(), paths=ok) == ok
    d = w.export_progress
    assert d.isVisible() and d.state == "ok" and "tugadi" in d.status.text()
    assert d.b_close.text() == "Yopish" and not d.b_retry.isVisible() and d.list.count() == 2

    missing = tmp_path / "yoq"                                # papka yo'q - ikkinchi fayl xato
    paths = [str(tmp_path / "c.svg"), str(missing / "d.svg")]
    assert w.run_export("svg", b.floors, RenderOptions(), paths=paths) == paths[:1]
    d = w.export_progress
    assert d.state == "partial" and "1 tasida xato" in d.status.text()
    assert d.failed()[0].error.startswith("Papka topilmadi") and d.b_retry.isVisible()
    missing.mkdir()
    assert d.run() == paths and d.state == "ok"              # qayta urinish - faqat xatolisi qayta yoziladi

    assert w.run_export("pdf", b.floors, RenderOptions(), paths=[str(tmp_path / "x" / "y.pdf")]) == []
    assert w.export_progress.state == "failed" and "bajarilmadi" in w.export_progress.status.text()
    assert "bajarilmadi" in w.statusBar().currentMessage()

    xl = w._write_excel(str(tmp_path / "h"), [b])
    assert xl.endswith("h.xlsx") and w.export_progress.state == "ok"
    w.export_progress.close()
    _close(w)


def test_import_screen_with_pdf(qapp, tmp_path):
    from PySide6.QtCore import QMarginsF, QRectF
    from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
    from PySide6.QtPdf import QPdfDocument
    from ttysi_cadastre.ui.import_dialog import ImportDialog, cv_to_qimage
    p = str(tmp_path / "k.pdf")
    wr = QPdfWriter(p)
    wr.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    painter = QPainter(wr)
    r = wr.pageLayout().paintRectPixels(wr.resolution())
    painter.drawImage(QRectF(0, 0, r.width(), r.height()), cv_to_qimage(make_plan(False)))
    wr.newPage()
    painter.drawImage(QRectF(0, 0, r.width(), r.height()), cv_to_qimage(make_plan(False)))
    painter.end()
    doc = QPdfDocument()
    doc.load(p)

    w = _window(Building(name="PDF"))
    dlg = ImportDialog(None, "1-qavat", w, doc=doc, page=1, file_name="k.pdf")
    assert dlg.page_list.count() == 2 and dlg.page() == 1
    assert dlg.img.shape[1] == pytest.approx(420 / 25.4 * 200, rel=0.01)
    dlg.scale.setValue(100)
    assert dlg.m_per_px() == pytest.approx(25.4 / 200 * 100 / 1000)
    assert "Masshtab hisoblandi" in dlg.scale_hint.text()
    assert "1-qavat" in dlg.size_tag.text() and " m" in dlg.size_tag.text()
    dlg.dpi.setCurrentIndex(0)                                # 150 DPI - sahifa qayta chiziladi
    assert dlg.img.shape[1] == pytest.approx(420 / 25.4 * 150, rel=0.01)

    # "Yaratish va yana etaj": oyna yopilganda etaj qo'shiladi va again=True
    from PySide6.QtCore import QTimer
    QTimer.singleShot(0, dlg.accept_again)
    assert w.create_floor(dlg)
    assert dlg.again
    f = w.building.floors[-1]
    assert f.scale_m_per_px and len(f.rooms) == len(ROOMS)
    assert w.floor is f and w.pages.currentWidget() is w.editor
    _close(w)
    doc.close()


def test_buildings_folder_and_floors(qapp, tmp_path):
    """Binolar papkasi: 1-bino > 1,2-qavat; 2-bino > 1-qavat. Bosh sahifada etaj bosilsa - shu etaj ochiladi."""
    from ttysi_cadastre.ui.home import ProjectCard, WORKSPACE_KEY, forget, scan
    w = _window()
    w.settings.setValue(WORKSPACE_KEY, str(tmp_path / "Binolar"))
    w.settings.remove("recent/projects")
    p1 = w.create_building("1-bino", "Toshkent")
    assert p1 == str(tmp_path / "Binolar" / "1-bino.ttysi") and os.path.exists(p1)
    for n in (1, 2):
        w.building.floors.append(Floor(name=f"{n}-qavat", width=100, height=100))
    w.building.floors[1].equipment.append(Equipment(code="?", x=5, y=5))
    assert w.save_project()
    w.create_building("2-bino")
    w.building.floors.append(Floor(name="1-qavat", width=100, height=100))
    w.save_project()
    assert w.create_building("1-bino").endswith("1-bino (2).ttysi")     # nom band - yangi fayl
    forget(w.settings, w.path)
    os.remove(w.path)

    items = scan(str(tmp_path / "Binolar"))
    assert [i["name"] for i in items] == ["1-bino", "2-bino"]
    assert items[0]["floor_list"] == [["1-qavat", 0], ["2-qavat", 1]]

    w.dirty = False
    w.show_home()
    cards = [c for c in w.home.cards if isinstance(c, ProjectCard)]
    assert len(cards) == 2 and not w.home.other_title.isVisible()
    chips = cards[0].findChildren(type(w.home.cards[-1]).__mro__[1], "floorChip")
    assert [c.text() for c in chips] == ["1-qavat", "2-qavat  ⚠1"]
    chips[1].click()                                                  # 1-bino, 2-qavat
    assert w.pages.currentWidget() is w.editor
    assert w.building.name == "1-bino" and w.floor is w.building.floors[1]

    w.fill_switch_menu()
    names = [a.text() for a in w.switch_menu.actions() if a.text()]
    assert names[0].startswith("1-bino") and w.switch_menu.actions()[0].isChecked()
    assert names[1].startswith("2-bino")
    _close(w)


def test_organizations(qapp, tmp_path):
    """Tashkilot: bir nechta binoni birlashtirish, ko'chirish, nomini o'zgartirish, hisobot, tarqatish."""
    from openpyxl import load_workbook
    from ttysi_cadastre.ui.home import ProjectCard, WORKSPACE_KEY, org_of, orgs, recent, scan
    ws = tmp_path / "Binolar"
    w = _window()
    w.settings.setValue(WORKSPACE_KEY, str(ws))
    w.settings.remove("recent/projects")
    paths = []
    for n in (1, 2, 3):
        paths.append(w.create_building(f"{n}-bino"))
        w.building.floors.append(Floor(name="1-qavat", width=100, height=100))
        w.building.floors[0].equipment.append(Equipment(code="K", x=5, y=5))
        w.save_project()
    w.dirty = False
    w.show_home()
    h = w.home

    org = h.create_org("5-maktab", paths[:2])                # 1 va 2-bino birlashtirildi
    assert os.path.isdir(org) and os.path.basename(org) == "5-maktab"
    o = orgs(str(ws))
    assert [x["name"] for x in o] == ["5-maktab"]
    assert [b["name"] for b in o[0]["buildings"]] == ["1-bino", "2-bino"]
    assert [b["name"] for b in scan(str(ws))] == ["3-bino"]
    assert w.path == os.path.join(org, "3-bino.ttysi") or w.path == paths[2]   # ochiq bino (3) joyida
    assert all(os.path.dirname(i["path"]) in (org, str(ws)) for i in recent(w.settings))
    cards = [c for c in h.cards if isinstance(c, ProjectCard)]
    assert len(cards) == 3 and not h.other_title.isVisible()

    p3 = h.move(paths[2], org)                                # 3-bino ham qo'shildi
    assert os.path.dirname(p3) == org and w.path == p3     # ochiq loyiha yo'li yangilandi
    assert org_of(w.settings, p3) == org
    w.open_project(p3)
    assert w.project_name.text() == "5-maktab › 3-bino"
    w.show_home()

    new_org = h.rename_org(org, "5-maktab (yangi)")
    assert os.path.basename(new_org) == "5-maktab (yangi)" and w.path == os.path.join(new_org, "3-bino.ttysi")

    out = w.export_excel_org(new_org, str(tmp_path / "hisobot.xlsx"))
    rows = list(load_workbook(out)["Umumiy"].iter_rows(values_only=True))
    assert {r[0] for r in rows[1:-1]} == {"1-bino", "2-bino", "3-bino"}
    assert rows[-1][0] == "5-MAKTAB (YANGI) - JAMI" and 3 in rows[-1]

    w.create_building("4-bino", org=new_org)                  # yangi bino tashkilot ichida
    assert os.path.dirname(w.path) == new_org
    w.dirty = False
    w.fill_switch_menu()
    sub = [a for a in w.switch_menu.actions() if a.menu()]
    assert sub and sub[0].text().startswith("5-maktab (yangi)")

    h.dissolve_org(new_org, confirm=False)                    # binolar o'chmaydi
    assert not os.path.exists(new_org) and orgs(str(ws)) == []
    assert [b["name"] for b in scan(str(ws))] == ["1-bino", "2-bino", "3-bino", "4-bino"]
    assert os.path.dirname(w.path) == str(ws)
    _close(w)


def test_organization_full_export(qapp, tmp_path):
    """Tashkilotni to'liq eksport: Excel, umumiy va har bino PDF, PNG/DXF, loyiha fayllari, ZIP."""
    import zipfile
    from PySide6.QtWidgets import QDialog
    from ttysi_cadastre.ui.home import WORKSPACE_KEY
    w = _window()
    w.settings.setValue(WORKSPACE_KEY, str(tmp_path / "Binolar"))
    w.settings.remove("recent/projects")
    paths = []
    for n in (1, 2):
        paths.append(w.create_building(f"{n}-bino"))
        for k in (1, 2):
            f = Floor(name=f"{k}-qavat", width=400, height=300, scale_m_per_px=0.05)
            f.rooms = [Room(points=[(20, 20), (200, 20), (200, 200), (20, 200)], name="1")]
            f.equipment = [Equipment(code="K", x=50, y=50)]
            w.building.floors.append(f)
        w.save_project()
    w.dirty = False
    w.show_home()
    org = w.home.create_org("5-maktab", paths)
    w.open_project(os.path.join(org, "2-bino.ttysi"))
    w.building.floors[0].equipment.append(Equipment(code="T", x=80, y=80))   # saqlanmagan o'zgarish
    w.dirty = True
    out = tmp_path / "eksport"
    out.mkdir()

    def hook(dlg):
        assert dlg.ok.isEnabled()
        dlg.dest.setText(str(out))
        for k in ("png", "dxf", "zip"):
            dlg.c[k].setChecked(True)
        return QDialog.Accepted
    saved = w.export_org_full(org, dialog_hook=hook)
    d = w.export_progress
    assert d.state == "ok", [j.error for j in d.jobs]
    folder = [p for p in out.iterdir() if p.is_dir()][0]
    names = {str(p.relative_to(folder)) for p in folder.rglob("*") if p.is_file()}
    assert {"5-maktab - hisobot.xlsx", "5-maktab - barcha binolar.pdf", "Tarkib.txt",
            os.path.join("1-bino", "1-bino.pdf"), os.path.join("2-bino", "2-bino - 2-qavat.png"),
            os.path.join("1-bino", "1-bino - 1-qavat.dxf"),
            os.path.join("Loyiha fayllari", "2-bino.ttysi")} <= names
    assert len(saved) == len(d.jobs) == 1 + 1 + 2 + 4 + 4 + 2 + 1 + 1
    copy = Building.load(str(folder / "Loyiha fayllari" / "2-bino.ttysi"))
    assert any(e.code == "T" for e in copy.floors[0].equipment)        # ochiq binoning o'zgarishlari ham
    txt = (folder / "Tarkib.txt").read_text(encoding="utf-8")
    assert "1-bino: 2 etaj" in txt and "hisobot.xlsx" in txt
    z = zipfile.ZipFile(str(folder) + ".zip")
    assert any(n.endswith("barcha binolar.pdf") for n in z.namelist())
    _close(w)
