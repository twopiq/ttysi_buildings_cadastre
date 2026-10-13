import os

import ezdxf
import pytest

from tests.synthetic import MARKS, ROOMS, make_double_line_plan, make_phone_photo, make_plan
from ttysi_cadastre.detection import DetectParams, detect, find_paper_corners, warp_to_rect
from ttysi_cadastre.geometry import merge_rooms, split_room
from ttysi_cadastre.model import Building, Equipment, Floor, Room


def _bbox(points):
    xs, ys = [p[0] for p in points], [p[1] for p in points]
    return min(xs), min(ys), max(xs), max(ys)


def test_detect_scanned_plan():
    res = detect(make_plan())
    assert len(res.rooms) == len(ROOMS)
    found = sorted(_bbox(r.points) for r in res.rooms)
    for exp, got in zip(sorted(ROOMS), found):
        assert all(abs(a - b) < 8 for a, b in zip(exp, got)), (exp, got)
    kinds = [r.kind for r in res.rooms]
    assert kinds.count("corridor") == 1
    assert len(res.marks) == len(MARKS)


def test_detect_phone_photo():
    photo = make_phone_photo(make_plan())
    img = warp_to_rect(photo, find_paper_corners(photo))
    res = detect(img)
    assert len(res.rooms) == len(ROOMS)
    assert len(res.marks) == len(MARKS)


def test_detect_double_line_walls_with_scale():
    """Kadastr uslubi: qo'sh ingichka chiziqli devorlar, eshik yoylari, masshtab ma'lum."""
    mpp = 0.025
    res = detect(make_double_line_plan(mpp), DetectParams(m_per_px=mpp, wall_m=0.3, door_m=1.2))
    assert len(res.rooms) == len(ROOMS)
    found = sorted(_bbox(r.points) for r in res.rooms)
    for exp, got in zip(sorted(ROOMS), found):
        assert all(abs(a - b) < 10 for a, b in zip(exp, got)), (exp, got)
    assert [r.kind for r in res.rooms].count("corridor") == 1


def test_corridor_with_exterior_entrance_is_kept():
    """Tashqi kirish eshigi bor koridor tashqi maydon deb tushib qolmasligi kerak."""
    mpp = 0.025
    res = detect(make_double_line_plan(mpp, entrance=True), DetectParams(m_per_px=mpp, wall_m=0.3, door_m=1.2))
    corridors = [r for r in res.rooms if r.kind == "corridor"]
    assert len(corridors) == 1
    x0, y0, x1, y1 = _bbox(corridors[0].points)
    assert x1 - x0 > 900 and 100 < y1 - y0 < 140
    assert len(res.rooms) == len(ROOMS)


def test_room_contains_and_room_of():
    f = Floor(name="x")
    big = Room(points=[(0, 0), (100, 0), (100, 100), (0, 100)], name="katta")
    small = Room(points=[(10, 10), (30, 10), (30, 30), (10, 30)], name="kichik")
    f.rooms = [big, small]
    assert f.room_of(Equipment(code="T", x=20, y=20)) is small
    assert f.room_of(Equipment(code="T", x=60, y=60)) is big
    assert f.room_of(Equipment(code="T", x=160, y=60)) is None


def test_split_and_merge():
    sq = [(0, 0), (100, 0), (100, 50), (0, 50)]
    parts = split_room(sq, (50, 10), (50, 40))
    assert len(parts) == 2
    assert merge_rooms(parts, 2) is not None
    assert merge_rooms([sq, [(300, 0), (400, 0), (400, 50), (300, 50)]], 5) is None


def test_project_roundtrip(tmp_path):
    b = Building(name="Bino", address="Manzil")
    f = Floor(name="1-etaj", image=b"\x89PNGfake", width=100, height=80, scale_m_per_px=0.05)
    f.rooms.append(Room(points=[(0, 0), (10, 0), (10, 10)], kind="corridor", name="K"))
    f.equipment.append(Equipment(code="W", x=5, y=5, note="izoh"))
    b.floors.append(f)
    p = tmp_path / "t.ttysi"
    b.save(str(p))
    b2 = Building.load(str(p))
    f2 = b2.floors[0]
    assert (b2.name, b2.address) == ("Bino", "Manzil")
    assert f2.image == f.image and f2.scale_m_per_px == 0.05
    assert f2.rooms[0].kind == "corridor" and f2.equipment[0].note == "izoh"


@pytest.fixture(scope="module")
def qapp():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _sample():
    b = Building(name="Bino")
    f = Floor(name="1-etaj", width=1000, height=700, scale_m_per_px=0.02)
    f.rooms = [Room(points=[(100, 100), (500, 100), (500, 400), (100, 400)], name="101"),
               Room(points=[(500, 100), (900, 100), (900, 400), (500, 400)], kind="corridor")]
    f.equipment = [Equipment(code=c, x=200 + 80 * i, y=250) for i, c in enumerate(["SER", "KX", "K", "W", "S", "T", "?"])]
    b.floors.append(f)
    return b, f


def test_exports(qapp, tmp_path):
    from ttysi_cadastre.export import export_dxf, export_image, export_pdf, export_svg
    from ttysi_cadastre.render import RenderOptions
    b, f = _sample()
    for name, fn in (("a.png", lambda p: export_image(p, b, f, RenderOptions(), 1000)),
                     ("a.jpg", lambda p: export_image(p, b, f, RenderOptions(), 1000)),
                     ("a.svg", lambda p: export_svg(p, b, f, RenderOptions())),
                     ("a.pdf", lambda p: export_pdf(p, b, [f, f], RenderOptions())),
                     ("a.dxf", lambda p: export_dxf(p, b, f))):
        p = str(tmp_path / name)
        fn(p)
        assert os.path.getsize(p) > 500, name
    doc = ezdxf.readfile(str(tmp_path / "a.dxf"))
    assert not doc.audit().has_errors
    msp = doc.modelspace()
    assert len(msp.query("LWPOLYLINE")) == 2
    assert len(msp.query("INSERT")) == 7
    assert {e.dxf.layer for e in msp.query("INSERT")} >= {"AT_W", "AT_TASDIQLANMAGAN"}


def test_main_window(qapp):
    from ttysi_cadastre.ui.main_window import MainWindow
    b, f = _sample()
    w = MainWindow()
    w.building = b
    w.refresh_all()
    assert w.floor is f
    # '?' belgini tanlab, 4 (Wi-Fi) tugmasi bilan tasdiqlash
    w.select_next_unknown()
    assert w.on_digit("4")
    assert f.equipment[-1].code == "W"
    w.undo()
    assert f.equipment[-1].code == "?"
    w.dirty = False
    w.close()


def test_excel_report(tmp_path):
    import openpyxl
    from ttysi_cadastre.report import export_excel
    b, f = _sample()
    f.equipment += [Equipment(code="T", x=150, y=150), Equipment(code="T", x=160, y=150)]
    p = str(tmp_path / "h.xlsx")
    export_excel(p, [b, b])
    wb = openpyxl.load_workbook(p)
    rows = list(wb["Umumiy"].iter_rows(values_only=True))
    head = rows[0]
    first = dict(zip(head, rows[1]))
    assert first["T"] == 3 and first["W"] == 1 and first["Tasdiqlanmagan"] == 1
    total = dict(zip(head, rows[-1]))
    assert total["Bino"] == "BARCHA BINOLAR" and total["T"] == 6
    rooms = [dict(zip(rows_h, r)) for rows_h in [next(wb["Xonalar"].iter_rows(values_only=True))]
             for r in list(wb["Xonalar"].iter_rows(values_only=True))[1:]]
    r101 = next(r for r in rooms if r["Xona №"] == "101")
    assert r101["T"] == 2 and r101["W"] == 1 and r101["Maydon, m²"] == 48.0
    corridor = next(r for r in rooms if r["Turi"] == "Koridor")
    assert corridor["T"] == 1 and corridor["Tasdiqlanmagan"] == 1
    assert wb["Jihozlar"].max_row == 1 + 2 * 9


def test_pdf_page_render(qapp, tmp_path):
    """PDF sahifasini rasm va qog'oz o'lchami (mm/piksel) bilan olish."""
    from PySide6.QtCore import QMarginsF, QRectF
    from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
    from PySide6.QtPdf import QPdfDocument
    from ttysi_cadastre.ui.import_dialog import PdfPageDialog, cv_to_qimage
    p = str(tmp_path / "k.pdf")
    w = QPdfWriter(p)
    w.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    painter = QPainter(w)
    r = w.pageLayout().paintRectPixels(w.resolution())
    painter.drawImage(QRectF(0, 0, r.width(), r.height()), cv_to_qimage(make_plan(False)))
    painter.end()
    doc = QPdfDocument()
    doc.load(p)
    dlg = PdfPageDialog(doc)
    img, mm_per_px = dlg.render()
    assert img.shape[1] == pytest.approx(420 / 25.4 * 200, rel=0.01)
    assert mm_per_px == pytest.approx(25.4 / 200)
    assert len(detect(img).rooms) == len(ROOMS)


def test_planned_equipment(qapp, tmp_path):
    """Rejadagi (kerak) jihozlar: saqlash, Excel, DXF va oynadagi almashtirish."""
    import openpyxl
    from ttysi_cadastre.export import export_dxf, export_image
    from ttysi_cadastre.render import RenderOptions
    from ttysi_cadastre.report import export_excel
    b, f = _sample()
    f.rooms[0].kind = "computer"
    f.equipment += [Equipment(code="S", x=150, y=150, status="planned", note="kompyuter xona uchun"),
                    Equipment(code="T", x=160, y=150, status="planned")]

    # fayl
    p = str(tmp_path / "r.ttysi")
    b.save(p)
    f2 = Building.load(p).floors[0]
    assert [e.status for e in f2.equipment].count("planned") == 2

    # Excel: umumiy varaqda faqat mavjudlar, "Kerak (reja)" varag'ida rejadagilar
    x = str(tmp_path / "r.xlsx")
    export_excel(x, [b])
    wb = openpyxl.load_workbook(x)
    rows = list(wb["Umumiy"].iter_rows(values_only=True))
    first = dict(zip(rows[0], rows[1]))
    assert first["S"] == 1 and first["T"] == 1 and first["Kerak (rejada)"] == 2
    plan = list(wb["Kerak (reja)"].iter_rows(values_only=True))
    floor_row = dict(zip(plan[0], plan[1]))
    assert floor_row["S"] == 1 and floor_row["T"] == 1 and floor_row["Jami"] == 2
    assert any(r and r[2] == "101" and r[3] == "S" for r in plan)
    rooms = list(wb["Xonalar"].iter_rows(values_only=True))
    r101 = dict(zip(rooms[0], next(r for r in rooms if r[2] == "101")))
    assert r101["Turi"] == "Kompyuter xona" and r101["Kerak (rejada)"] == "S×1, T×1"
    jih = list(wb["Jihozlar"].iter_rows(values_only=True))
    assert [r[6] for r in jih[1:]].count("Kerak") == 2

    # DXF: alohida punktir qatlam
    d = str(tmp_path / "r.dxf")
    export_dxf(d, b, f)
    doc = ezdxf.readfile(d)
    assert doc.layers.get("AT_S_KERAK").dxf.linetype == "DASHED"
    assert {e.dxf.layer for e in doc.modelspace().query("INSERT")} >= {"AT_S_KERAK", "AT_T_KERAK", "AT_S"}
    export_image(str(tmp_path / "r.png"), b, f, RenderOptions(), 1000)

    # oyna: Ctrl+K bilan holatni almashtirish va bekor qilish
    from ttysi_cadastre.ui.main_window import MainWindow
    w = MainWindow()
    w.building = b
    w.refresh_all()
    item = next(i for i in w.scene.equipment_items() if i.eq.code == "W")
    item.setSelected(True)
    w.toggle_status()
    assert item.eq.planned
    w.undo()
    assert not next(e for e in f.equipment if e.code == "W").planned
    w.dirty = False
    w.close()


def test_doors_and_stairs(qapp, tmp_path):
    """Eshik devorga yopishadi; eshik va zina saqlanadi, chiziladi va DXF'ga chiqadi."""
    from PySide6.QtCore import QPointF
    from ttysi_cadastre.export import export_dxf, export_image
    from ttysi_cadastre.geometry import door_lines, stair_lines
    from ttysi_cadastre.model import Door
    from ttysi_cadastre.render import RenderOptions
    from ttysi_cadastre.ui.main_window import MainWindow
    b, f = _sample()
    f.rooms.append(Room(points=[(100, 400), (200, 400), (200, 600), (100, 600)], kind="stairs"))
    w = MainWindow()
    w.building = b
    w.refresh_all()
    # vertikal devor (x=500) yoniga bosish -> eshik devorga yopishadi va vertikal bo'ladi
    item = w.scene.add_door(QPointF(507, 250))
    d = item.door
    assert abs(d.x - 500) < 1e-6 and abs(abs(d.angle) - 90) < 1e-6
    assert abs(d.width - 0.9 / f.scale_m_per_px) < 1e-6
    assert len(f.doors) == 1
    w.undo()
    assert not f.doors
    w.redo()
    assert len(f.doors) == 1
    f.doors[0].leaves = 2
    # bir tabaqali: 2 kesak + tabaqa + yoy; ikki tabaqali: 2 + 2*2
    assert len(door_lines(Door(x=0, y=0, width=40), 5)) == 4
    assert len(door_lines(f.doors[0], 5)) == 6
    steps = stair_lines(f.rooms[-1].points, 0.3 / f.scale_m_per_px)
    assert len(steps) > 10

    p = str(tmp_path / "e.ttysi")
    b.save(p)
    f2 = Building.load(p).floors[0]
    assert f2.doors[0].leaves == 2 and f2.rooms[-1].kind == "stairs"

    export_image(str(tmp_path / "e.png"), b, f, RenderOptions(), 1000)
    dx = str(tmp_path / "e.dxf")
    export_dxf(dx, b, f)
    msp = ezdxf.readfile(dx).modelspace()
    assert len(msp.query('LWPOLYLINE[layer=="ESHIKLAR"]')) == 6
    assert len(msp.query('LWPOLYLINE[layer=="ZINALAR"]')) > 10
    w.dirty = False
    w.close()


def test_rotate_in_place(qapp, monkeypatch):
    """Xona, zina va eshik o'z markazi atrofida aylanadi; bekor qilish ishlaydi."""
    from PySide6.QtCore import QPointF
    from ttysi_cadastre.ui import canvas
    from ttysi_cadastre.ui.main_window import MainWindow
    b, f = _sample()
    room = f.rooms[0]                                   # 400x300, markazi (300, 250)
    c0, a0 = room.centroid(), room.area()
    w = MainWindow()
    w.building = b
    w.refresh_all()
    item = next(i for i in w.scene.room_items() if i.room is room)
    item.setSelected(True)
    w.rotate_selected(90.0)
    xs, ys = [p[0] for p in room.points], [p[1] for p in room.points]
    assert (max(xs) - min(xs), max(ys) - min(ys)) == (300, 400)      # eni va bo'yi almashdi
    assert room.centroid() == pytest.approx(c0) and room.area() == pytest.approx(a0)
    w.undo()
    room = f.rooms[0]
    assert sorted(room.points) == sorted([(100, 100), (500, 100), (500, 400), (100, 400)])

    # eshik: 180° - ochilish tomoni teskari, joyi o'zgarmaydi
    d = w.scene.add_door(QPointF(300, 102)).door
    x, y, a = d.x, d.y, d.angle
    w.scene.clearSelection()
    door_item = next(i for i in w.scene.items() if isinstance(i, canvas.DoorItem))
    door_item.setSelected(True)
    w.rotate_selected(180.0)
    assert (d.x, d.y) == (x, y) and d.angle == pytest.approx((a + 180) % 360)

    # istalgan burchak (dialog o'rniga 30°)
    monkeypatch.setattr(canvas, "ask_angle", lambda parent=None: 30.0)
    w.rotate_selected("custom")
    assert d.angle == pytest.approx((a + 210) % 360)

    # hech narsa tanlanmagan - o'zgarish yo'q
    w.scene.clearSelection()
    assert not w.scene.rotate_selected(90.0)
    w.dirty = False
    w.close()


def test_walls(qapp, tmp_path):
    """Mustaqil devor: asbob bilan chizish, eshik yopishishi, aylantirish, saqlash, DXF."""
    from PySide6.QtCore import QPointF, Qt
    from PySide6.QtTest import QTest
    from ttysi_cadastre.export import export_dxf, export_image
    from ttysi_cadastre.render import RenderOptions, door_wall_thickness
    from ttysi_cadastre.ui import canvas
    from ttysi_cadastre.ui.main_window import MainWindow
    b, f = _sample()                                     # 1 m = 50 px
    w = MainWindow()
    w.resize(1200, 800)
    w.show()
    w.building = b
    w.refresh_all()
    qapp.processEvents()
    w.view.zoom_fit()
    qapp.processEvents()

    # asbob bilan: 3 nuqta + Enter -> siniq devor (orto rejimda)
    w.select_tool(canvas.WALL)
    vp = w.view.viewport()
    for x, y in [(150, 500), (450, 503), (450, 650)]:
        QTest.mouseClick(vp, Qt.LeftButton, Qt.NoModifier, w.view.mapFromScene(QPointF(x, y)))
    QTest.keyClick(w.view, Qt.Key_Return)
    assert len(f.walls) == 1
    wall = f.walls[0]
    assert len(wall.points) == 3
    (x0, y0), (x1, y1), (x2, y2) = wall.points
    assert abs(y1 - y0) < 1e-6 and abs(x2 - x1) < 1e-6          # orto: gorizontal, keyin vertikal
    assert wall.thickness == pytest.approx(0.25 / f.scale_m_per_px)

    # eshik mustaqil devorga yopishadi va uning qalinligini oladi
    wall.thickness = 0.5 / f.scale_m_per_px
    door = w.scene.add_door(QPointF(300, y0 + 6)).door
    assert abs(door.y - y0) < 1e-6 and abs(door.angle) < 1e-6
    assert door_wall_thickness(f, door) == pytest.approx(wall.thickness)

    # aylantirish va bekor qilish
    w.select_tool(canvas.SELECT)
    w.scene.clearSelection()
    item = next(i for i in w.scene.items() if isinstance(i, canvas.WallItem))
    item.setSelected(True)
    before = list(f.walls[0].points)
    w.rotate_selected(90.0)
    assert f.walls[0].points != before
    w.undo()
    assert f.walls[0].points == pytest.approx(before)

    # saqlash, chizish, DXF
    p = str(tmp_path / "d.ttysi")
    b.save(p)
    assert len(Building.load(p).floors[0].walls[0].points) == 3
    export_image(str(tmp_path / "d.png"), b, f, RenderOptions(), 1000)
    dx = str(tmp_path / "d.dxf")
    export_dxf(dx, b, f)
    walls = ezdxf.readfile(dx).modelspace().query('LWPOLYLINE[layer=="DEVORLAR"]')
    assert len(walls) == 1 and walls[0].dxf.const_width == pytest.approx(0.5)
    w.dirty = False
    w.close()
