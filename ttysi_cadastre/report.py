"""Excel hisobot: bino(lar) bo'yicha AT jihozlari soni - etajlar, xonalar va jihozlar ro'yxati."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .equipment import EQUIPMENT_TYPES, ROOM_KINDS, UNKNOWN, get_type
from .model import Building, Floor

CODES = [t.code for t in EQUIPMENT_TYPES] + [UNKNOWN.code]
HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _type_headers() -> list[str]:
    return [t.label for t in EQUIPMENT_TYPES] + ["Tasdiqlanmagan"]


def _counts(items) -> list[int]:
    c = {code: 0 for code in CODES}
    for e in items:
        c[e.code if e.code in c else UNKNOWN.code] += 1
    return [c[code] for code in CODES]


def _need_text(items) -> str:
    """Rejadagi jihozlar qisqacha: "S×1, T×2"."""
    c = {}
    for e in items:
        if e.planned:
            c[e.code] = c.get(e.code, 0) + 1
    return ", ".join(f"{get_type(k).label}×{v}" for k, v in c.items())


def _area(floor: Floor, room) -> float | None:
    return round(room.area() * floor.scale_m_per_px ** 2, 2) if floor.scale_m_per_px else None


def _header(ws, row: list[str]):
    ws.append(row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _row(ws, row: list, total: bool = False):
    ws.append(row)
    for cell in ws[ws.max_row]:
        cell.border = BORDER
        if total:
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL


def _widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def export_excel(path: str, buildings: list[Building]) -> None:
    wb = Workbook()

    # --- 1. Umumiy: bino -> etaj -> jihozlar soni
    ws = wb.active
    ws.title = "Umumiy"
    _header(ws, ["Bino", "Etaj"] + _type_headers() + ["Xonalar soni", "Maydon, m²", "Kerak (rejada)"])
    grand = [0] * len(CODES)
    for b in buildings:
        sub = [0] * len(CODES)
        for f in b.floors:
            c = _counts(e for e in f.equipment if not e.planned)
            area = round(sum(r.area() for r in f.rooms) * f.scale_m_per_px ** 2, 1) if f.scale_m_per_px else None
            _row(ws, [b.name, f.name] + c + [len(f.rooms), area, sum(e.planned for e in f.equipment)])
            sub = [a + x for a, x in zip(sub, c)]
        if len(buildings) > 1 or len(b.floors) > 1:
            _row(ws, [b.name, "Jami"] + sub + [sum(len(f.rooms) for f in b.floors), None,
                                               sum(e.planned for f in b.floors for e in f.equipment)], total=True)
        grand = [a + x for a, x in zip(grand, sub)]
    if len(buildings) > 1:
        _row(ws, ["BARCHA BINOLAR", ""] + grand + [None, None,
                  sum(e.planned for b in buildings for f in b.floors for e in f.equipment)], total=True)
    _widths(ws, [28, 14] + [9] * (len(CODES) - 1) + [15, 12, 12, 14])

    # --- 2. Xonalar: har bir xonadagi jihozlar
    ws = wb.create_sheet("Xonalar")
    _header(ws, ["Bino", "Etaj", "Xona №", "Turi", "Maydon, m²"] + _type_headers() + ["Kerak (rejada)", "Izohlar"])
    for b in buildings:
        for f in b.floors:
            by_room: dict[str, list] = {}
            outside = []
            for e in f.equipment:
                r = f.room_of(e)
                (by_room.setdefault(r.id, []) if r else outside).append(e)
            for r in f.rooms:
                items = by_room.get(r.id, [])
                notes = "; ".join(f"{get_type(e.code).label} {e.note}" for e in items if e.note)
                _row(ws, [b.name, f.name, r.name, ROOM_KINDS.get(r.kind, ("",))[0], _area(f, r)]
                     + _counts(e for e in items if not e.planned) + [_need_text(items), notes])
            if outside:
                _row(ws, [b.name, f.name, "(xonadan tashqarida)", "", None]
                     + _counts(e for e in outside if not e.planned) + [_need_text(outside), ""])
    _widths(ws, [28, 14, 12, 16, 11] + [9] * (len(CODES) - 1) + [15, 22, 40])
    ws.auto_filter.ref = ws.dimensions

    # --- 3. Jihozlar ro'yxati
    ws = wb.create_sheet("Jihozlar")
    _header(ws, ["№", "Bino", "Etaj", "Xona №", "Belgi", "Jihoz", "Holati", "Izoh"])
    n = 0
    for b in buildings:
        for f in b.floors:
            for e in sorted(f.equipment, key=lambda e: (CODES.index(e.code) if e.code in CODES else 99, e.y, e.x)):
                n += 1
                r = f.room_of(e)
                t = get_type(e.code)
                _row(ws, [n, b.name, f.name, r.name if r else "", t.label, t.name,
                          "Kerak" if e.planned else "Mavjud", e.note])
    _widths(ws, [6, 28, 14, 12, 8, 26, 10, 40])
    ws.auto_filter.ref = ws.dimensions

    # --- 4. Kerak (reja): nimani qayerga o'rnatish kerak
    ws = wb.create_sheet("Kerak (reja)")
    _header(ws, ["Bino", "Etaj"] + [t.label for t in EQUIPMENT_TYPES] + ["Jami"])
    plan_total = [0] * len(EQUIPMENT_TYPES)
    for b in buildings:
        for f in b.floors:
            c = _counts(e for e in f.equipment if e.planned)[:len(EQUIPMENT_TYPES)]
            if sum(c):
                _row(ws, [b.name, f.name] + c + [sum(c)])
                plan_total = [a + x for a, x in zip(plan_total, c)]
    _row(ws, ["JAMI", ""] + plan_total + [sum(plan_total)], total=True)
    ws.append([])
    _header(ws, ["Bino", "Etaj", "Xona №", "Belgi", "Jihoz", "Izoh"])
    for b in buildings:
        for f in b.floors:
            for e in f.equipment:
                if e.planned:
                    r = f.room_of(e)
                    t = get_type(e.code)
                    _row(ws, [b.name, f.name, r.name if r else "", t.label, t.name, e.note])
    _widths(ws, [28, 14] + [12] * len(EQUIPMENT_TYPES) + [10])

    wb.save(path)
